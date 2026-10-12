"""
Compare one wages period of PayWages (legacy Smart-Eye SQL Server) with paywages (tenant MySQL),
employee by employee, into an Excel workbook. Read-only on both databases.

    python dbqueries/compare_paywages_period.py 2026-09-01 2026-09-15 [out.xlsx]

Run from pyhback/. MySQL comes from env/database.env (TENANT_DB, default nbjcl). SQL Server:
MSSQL_SERVER (default 162.245.218.10, the live server), MSSQL_USER (default ERP) and MSSQL_PWD;
MSSQL_SERVER=localhost uses the Windows login and needs no password.

Rows are matched on ECode + PaySlipType for the period's end date. Audit columns that always
differ (created / updated stamps, PC and user names) are left out of the comparison.
"""
import datetime
import os
import sys
from collections import defaultdict
from decimal import Decimal

import pymysql
import pyodbc
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

IGNORED = ("CreatedDate", "UpdatedDate", "PCName", "UserID", "UserName")
# shown side by side for every employee
KEY_FIGURES = ("Work_HR", "NS_HR", "Holiday_HR", "FBasic_Amt", "DA_Amt", "Incentive_Amt", "PF_Gross",
               "Wages_Gross", "PFEmployee", "ESI_Amt", "PTax_Amt", "PFLoanCap", "Advance", "TotalDeduction",
               "NetPayable")


def norm(v):
    if v is None:
        return ""
    if isinstance(v, (Decimal, float, int)):
        return float(v)
    if isinstance(v, datetime.datetime):
        return v.replace(microsecond=0)
    if isinstance(v, str):
        return v.rstrip()
    return v


def same(x, y):
    if isinstance(x, float) and isinstance(y, float):
        return abs(x - y) <= 0.005
    if isinstance(x, datetime.datetime) and type(y) is datetime.date:
        return x.date() == y
    if isinstance(y, datetime.datetime) and type(x) is datetime.date:
        return y.date() == x
    return x == y


def main(from_date: str, to_date: str, out: str) -> None:
    load_dotenv("env/database.env")
    server = os.environ.get("MSSQL_SERVER", "162.245.218.10")
    if server.lower() in ("localhost", ".", "(local)"):
        auth = "Trusted_Connection=yes"
    else:
        if not os.environ.get("MSSQL_PWD"):
            sys.exit("Set MSSQL_PWD (SQL Server password for %s) or MSSQL_SERVER=localhost" % server)
        auth = "UID=%s;PWD=%s" % (os.environ.get("MSSQL_USER", "ERP"), os.environ["MSSQL_PWD"])
    ms = pyodbc.connect("DRIVER={ODBC Driver 17 for SQL Server};SERVER=%s;DATABASE=Smart_Eye_Jute_Northbrook_Live;"
                        "%s;TrustServerCertificate=yes" % (server, auth)).cursor()
    tenant = os.environ.get("TENANT_DB", "nbjcl")
    my = pymysql.connect(host=os.environ["DATABASE_HOST"], port=int(os.environ["DATABASE_PORT"]),
                         user=os.environ["DATABASE_USER"], password=os.environ["DATABASE_PASSWORD"],
                         database=tenant).cursor()

    ms_cols = [r[0] for r in ms.execute("select column_name from information_schema.columns "
                                        "where table_name = 'PayWages' order by ordinal_position")]
    my.execute("select column_name from information_schema.columns where table_schema = %s "
               "and table_name = 'paywages'", (tenant,))
    my_have = {r[0].lower() for r in my.fetchall()}
    cols = [c for c in ms_cols if c.lower() in my_have]
    cmp_cols = [c for c in cols if c not in IGNORED]
    ie, ip, iname = cols.index("ECode"), cols.index("PaySlipType"), cols.index("EName")

    # the period is the fortnight ending on to_date; from_date only documents it (PayWages has no from date)
    a_rows = ms.execute("select %s from PayWages where EDate = ? order by ECode, PaySlipType, SrlNo"
                        % ", ".join("[%s]" % c for c in cols), to_date).fetchall()
    my.execute("select %s from paywages where EDate = %%s order by ECode, PaySlipType, SrlNo"
               % ", ".join("`%s`" % c for c in cols), (to_date,))
    b_rows = my.fetchall()

    def keyed(rows):
        d, seen = {}, defaultdict(int)
        for r in rows:
            k = (str(r[ie]).strip(), str(r[ip]).strip())
            seen[k] += 1
            d[(*k, seen[k])] = r
        return d

    a, b = keyed(a_rows), keyed(b_rows)
    employees, cells, pairs = [], [], []
    col_count = defaultdict(int)
    counts = defaultdict(int)
    for k in sorted(set(a) | set(b)):
        ra, rb = a.get(k), b.get(k)
        name = str((ra or rb)[iname]).strip()
        fig = []
        for c in KEY_FIGURES:
            i = cols.index(c)
            x = norm(ra[i]) if ra else ""
            y = norm(rb[i]) if rb else ""
            fig += [x, y, (y - x) if isinstance(x, float) and isinstance(y, float) else ""]
        if ra is None or rb is None:
            status = "Only in MySQL" if ra is None else "Only in SQL Server"
            employees.append((k[0], name, k[1], status, "", "", *fig))
            counts[status] += 1
            pairs.append((ra, rb, None))
            continue
        diff = []
        for i, c in enumerate(cols):
            if c in IGNORED:
                continue
            x, y = norm(ra[i]), norm(rb[i])
            if not same(x, y):
                diff.append(c)
                col_count[c] += 1
                cells.append((k[0], name, k[1], c, x, y, (y - x) if isinstance(x, float) and isinstance(y, float) else ""))
        pairs.append((ra, rb, set(diff)))
        status = "Different" if diff else "Match"
        counts[status] += 1
        employees.append((k[0], name, k[1], status, len(diff) or "", ", ".join(diff), *fig))

    head_fill = PatternFill("solid", fgColor="1F4E78")
    bad, ok = PatternFill("solid", fgColor="FCE4D6"), PatternFill("solid", fgColor="E2EFDA")
    wb = Workbook()

    def sheet(ws, header, rows, widths, status_col=None):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for r in rows:
            ws.append(list(r))
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 14)
        if status_col is not None:
            for row in ws.iter_rows(min_row=2):
                row[status_col].fill = ok if row[status_col].value == "Match" else bad

    ws = wb.active
    ws.title = "Summary"
    for r in (
        ("Period", "%s to %s (rows with EDate = %s)" % (from_date, to_date, to_date)),
        ("Compared on", datetime.datetime.now().strftime("%d-%m-%Y %H:%M")),
        ("SQL Server", "Smart_Eye_Jute_Northbrook_Live.dbo.PayWages on %s" % server),
        ("MySQL", "%s.paywages on %s" % (tenant, os.environ["DATABASE_HOST"])),
        ("", ""),
        ("Rows in SQL Server", len(a_rows)),
        ("Rows in MySQL", len(b_rows)),
        ("Employees matching on every compared column", counts["Match"]),
        ("Employees with differences", counts["Different"]),
        ("Employees only in SQL Server", counts["Only in SQL Server"]),
        ("Employees only in MySQL", counts["Only in MySQL"]),
        ("Cells that differ", len(cells)),
        ("", ""),
        ("Columns compared", len(cmp_cols)),
        ("Columns left out", ", ".join(IGNORED)),
        ("Matching", "Rows are paired on ECode + PaySlipType. Numbers must agree to 0.005; text ignores trailing spaces."),
        ("Sheet 'Side by side'", "Two rows per employee: the first from SmartEye (SQL Server), the second from nbjcl "
                                 "(MySQL). A red cell differs between the two. A red Source cell with no red "
                                 "values means the employee exists on that side only."),
    ):
        ws.append(list(r))
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 46, 90
    for c in ws["A"]:
        c.font = Font(bold=True)

    # one pair of rows per employee: SmartEye first, nbjcl second; differing cells in red
    red_fill, red_font = PatternFill("solid", fgColor="FFC7CE"), Font(color="9C0006", bold=True)
    ws = wb.create_sheet("Side by side")
    ws.append(["Source", *cols])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head_fill
    for ra, rb, diff in pairs:
        for label, r in (("SmartEye", ra), ("nbjcl", rb)):
            if r is None:
                continue
            ws.append([label, *[norm(v) for v in r]])
            row = ws[ws.max_row]
            if diff is None:  # the employee exists on this side only
                row[0].fill, row[0].font = red_fill, red_font
                continue
            for i, c in enumerate(cols, 1):
                if c in diff:
                    row[i].fill, row[i].font = red_fill, red_font
            if diff:
                row[0].fill, row[0].font = red_fill, red_font
    ws.freeze_panes = ws.cell(row=2, column=cols.index("EName") + 3)
    ws.column_dimensions["A"].width = 11
    for i, c in enumerate(cols, 2):
        ws.column_dimensions[get_column_letter(i)].width = 26 if c in ("EName", "BenificiaryName") else 14

    fig_head = [h for c in KEY_FIGURES for h in (c + " SQL Server", c + " MySQL", c + " diff")]
    sheet(wb.create_sheet("Employees"),
          ["ECode", "Name", "PaySlipType", "Status", "Columns differing", "Which columns", *fig_head], employees,
          {"Name": 28, "Status": 18, "Which columns": 60, **{h: 16 for h in fig_head}}, status_col=3)
    sheet(wb.create_sheet("Cell differences"),
          ["ECode", "Name", "PaySlipType", "Column", "SQL Server value", "MySQL value", "Difference (MySQL - SQL Server)"],
          cells, {"Name": 28, "Column": 20, "SQL Server value": 26, "MySQL value": 26,
                  "Difference (MySQL - SQL Server)": 30})
    sheet(wb.create_sheet("Differences by column"), ["Column", "Employees differing"],
          sorted(col_count.items(), key=lambda x: -x[1]), {"Column": 24, "Employees differing": 20})
    wb.save(out)
    print("SQL Server rows %d | MySQL rows %d | %s | cells differing %d" % (
        len(a_rows), len(b_rows), dict(counts), len(cells)))
    print("saved", out)


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else
         "paywages_compare_%s.xlsx" % sys.argv[2])
